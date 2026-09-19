"""Offline regression tests: python -m unittest discover -s scripts -p test_*.py"""
import argparse
import tempfile
import unittest
from pathlib import Path
from docx import Document
from lxml import etree as E
import docx_ops as d


class DocumentTests(unittest.TestCase):
    def fixture(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'source.docx';doc=Document()
            p=doc.add_paragraph('正文引用')
            for s in ['[','1','-2',']，另见[2]。']:
                r=p.add_run(s);r.font.superscript=True
            doc.add_paragraph('参考文献')
            doc.add_paragraph('[1] First.');doc.add_paragraph('[2] Second.')
            doc.save(path);parts=d.read(path)
        parts['word/embeddings/test.bin']=b'unchanged-ole'
        return parts,E.fromstring(parts['word/document.xml'])

    def test_bibliography_split_runs(self):
        parts,root=self.fixture();original=d.text(root.xpath('//w:p',namespaces=d.NS)[0]);before=parts.copy()
        report=d.bibliography(parts,root,'参考文献')
        self.assertEqual(report['references'],2);self.assertEqual(report['REF_fields_added'],3)
        self.assertTrue(report['first_appearance_sequential'])
        self.assertEqual(d.text(root.xpath('//w:p',namespaces=d.NS)[0]),original)
        self.assertEqual([f['result'] for f in d.field_inventory(root)],['1','2','[2]'])
        self.assertEqual(len(root.xpath('//w:numPr',namespaces=d.NS)),2)
        self.assertTrue(all(r.find('w:rPr/w:vertAlign',d.NS) is not None for r in root.xpath('//w:r[w:instrText]',namespaces=d.NS)))
        self.assertTrue(d.audit_assets(before,parts))
        with self.assertRaises(ValueError):d.bibliography(parts,root,'参考文献')

    def test_protected_content(self):
        p=E.Element(d.W+'p');p.append(d.run('先[1]'))
        p.extend(d.field('REF existing \\h','[2]'));p.append(d.run('后[3]'))
        self.assertEqual([m[0] for m in d.CITE.finditer(d.stream(p)[0])],['[1]','[3]'])
        r=d.run('[4]');E.SubElement(r,d.W+'object');p.append(r)
        self.assertNotIn('[4]',d.stream(p)[0])

    def test_rendered_pagebreak(self):
        p=E.Element(d.W+'p');r=d.run('[1] Example');r.insert(0,E.Element(d.W+'lastRenderedPageBreak'));p.append(r)
        d.replace_match(p,0,4,lambda rp:[])
        self.assertEqual(d.text(p),'Example');self.assertEqual(len(p.xpath('.//w:lastRenderedPageBreak',namespaces=d.NS)),1)

    def test_normalize_opt_in(self):
        root=E.Element(d.W+'document');p=E.SubElement(root,d.W+'p');p.append(d.run('中文 AMZI 与 2 个 English words，2 μm。借鉴'))
        d.pprop(p,'spacing').set(d.W+'afterLines','100')
        cap=E.SubElement(root,d.W+'p');cap.append(d.run('图 1 流程'))
        args=argparse.Namespace(boundary_spaces=True,real_blank_lines=True,caption_fonts=True,adopt_wording=True)
        report=d.normalize(root,args)
        self.assertEqual(d.text(p),'中文AMZI与2个English words，2 μm。采用')
        self.assertEqual(d.text(cap),'图 1 流程');self.assertEqual(report['blank_paragraphs'],1)
        self.assertEqual(d.normalize(root,args)['blank_paragraphs'],0)

    def test_move_sentence_preserves_field(self):
        root=E.Element(d.W+'document')
        for s in ['方法','正文。具体研究流程如','理论']:
            p=E.SubElement(root,d.W+'p');p.append(d.run(s))
        p=root[1];p.extend(d.field('REF _RefFig1 \\h','图1'));p.append(d.run('所示。'))
        original_fields=d.field_inventory(root)
        args=argparse.Namespace(start_heading='方法',end_heading='理论',phrase='具体研究流程如',expected_count=1)
        d.lead_sentence(root,args)
        self.assertEqual(d.text(p),'具体研究流程如图1所示。正文。')
        self.assertEqual(d.field_inventory(root),original_fields)

    def test_figures(self):
        root=E.Element(d.W+'document');p=E.SubElement(root,d.W+'p');p.append(d.run('见图1。'))
        cap=E.SubElement(root,d.W+'p');cap.append(d.run('图'));cap.extend(d.field('SEQ 图 \\* ARABIC','1'));cap.append(d.run('流程'))
        self.assertEqual(d.figures(root)['REF_fields_added'],1)
        self.assertEqual(d.text(cap),'图 1 流程');self.assertEqual(d.text(p),'见图 1。')
        with self.assertRaises(ValueError):d.figures(root)

    def test_field_inventory_scoped_and_split_code(self):
        root=E.Element(d.W+'document');p=E.SubElement(root,d.W+'p');other=E.SubElement(root,d.W+'p')
        f=E.SubElement(other,d.W+'fldSimple');f.set(d.W+'instr','SEQ 图');f.append(d.run('1'))
        self.assertEqual(d.field_inventory(p),[])
        p.extend(d.field('REF target \\h','图 1'))
        t=p.find('.//w:instrText',d.NS);t.text=' REF '
        r=d.run();E.SubElement(r,d.W+'instrText').text='target \\h ';p.insert(2,r)
        self.assertEqual(d.field_inventory(p)[0]['code'],' REF target \\h ')
        self.assertEqual(len(d.field_inventory(root)),2)

    def test_simple_caption_and_body_size(self):
        root=E.Element(d.W+'document');p=E.SubElement(root,d.W+'p')
        rp=E.Element(d.W+'rPr');d.font_size(rp,12);p.append(d.run('见图1。',rp))
        cap=E.SubElement(root,d.W+'p');cap.append(d.run('图'))
        f=E.SubElement(cap,d.W+'fldSimple');f.set(d.W+'instr','SEQ 图 \\* ARABIC');f.append(d.run('1'))
        cap.append(d.run(' 流程'))
        d.figures(root,caption_size=14)
        self.assertEqual(d.text(p),'见图 1。')
        self.assertEqual(set(cap.xpath('./w:r/w:rPr/w:sz/@w:val',namespaces=d.NS)),{'28'})
        self.assertEqual(set(p.xpath('./w:r/w:rPr/w:sz/@w:val',namespaces=d.NS)),{'24'})

    def test_caption_fonts_preserve_position_and_field(self):
        root=E.Element(d.W+'document');p=E.SubElement(root,d.W+'p');p.append(d.run('图 '))
        f=E.SubElement(p,d.W+'fldSimple');f.set(d.W+'instr','SEQ 图');r=d.run('1');f.append(r)
        rp=E.Element(d.W+'rPr');r.insert(0,rp);E.SubElement(rp,d.W+'position').set(d.W+'val','-2')
        p.append(d.run(' 流程'))
        before=d.field_inventory(p)
        args=argparse.Namespace(boundary_spaces=True,real_blank_lines=False,caption_fonts=True,caption_size=14,adopt_wording=False)
        d.normalize(root,args)
        self.assertEqual(d.field_inventory(p),before)
        self.assertEqual(d.text(p),'图 1 流程')
        self.assertEqual(rp.find(d.W+'position').get(d.W+'val'),'-2')
        self.assertEqual(rp.find(d.W+'sz').get(d.W+'val'),'28')
        self.assertEqual(rp.find(d.W+'rFonts').get(d.W+'ascii'),'Times New Roman')

    def test_custom_caption_sequence_rejected(self):
        root=E.Element(d.W+'document');p=E.SubElement(root,d.W+'p');p.append(d.run('图'))
        p.extend(d.field('SEQ 图 \\r 3','3'));p.append(d.run(' 流程'))
        with self.assertRaises(ValueError):d.figures(root)

    def test_assets_compare_content_not_names(self):
        before={'word/embeddings/one.bin':b'ole','word/media/image1.png':b'image'}
        after={'word/embeddings/two.bin':b'ole','word/media/image2.png':b'image'}
        self.assertTrue(d.audit_assets(before,after))
        after['word/media/image2.png']=b'changed'
        with self.assertRaises(ValueError):d.audit_assets(before,after)

    def test_existing_output_refused(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'out.docx';d.save({'a':b'a'},p)
            with self.assertRaises(ValueError):d.save({'a':b'b'},p)

    def test_inherited_indent_and_number_tab(self):
        parts,root=self.fixture();ps=root.xpath('//w:p',namespaces=d.NS)
        styles=E.fromstring(parts['word/styles.xml'])
        style=E.SubElement(styles,d.W+'style');style.set(d.W+'styleId','ReferenceIndent')
        pp=E.SubElement(style,d.W+'pPr');ind=E.SubElement(pp,d.W+'ind')
        ind.set(d.W+'left','645');ind.set(d.W+'hanging','645');parts['word/styles.xml']=d.xml(styles)
        for p in ps[2:]:d.pprop(p,'pStyle').set(d.W+'val','ReferenceIndent')
        d.bibliography(parts,root,'参考文献')
        for p in ps[2:]:
            self.assertEqual(p.find('w:pPr/w:ind',d.NS).get(d.W+'left'),'645')
            self.assertEqual(p.find('w:pPr/w:tabs/w:tab',d.NS).get(d.W+'pos'),'645')


if __name__=='__main__':unittest.main()
